import codecs
import csv
import io
from datetime import datetime
from typing import Any, Iterable, Tuple
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models.entry import Entry

EXPORT_COLUMNS = [
    "Challan No",
    "Vehicle No",
    "Party",
    "Product",
    "Destination",
    "Quantity",
    "Unit Price",
    "Subtotal",
    "GST Type",
    "GST Rate (%)",
    "GST Amount",
    "Total Price",
    "Created At",
]


def sanitize_formula_injection(value: Any) -> Any:
    """
    Guards against CSV/Excel formula injection.
    Prefixes text fields starting with '=', '+', '-', or '@' with a single quote (').
    """
    if isinstance(value, str) and value:
        if value[0] in ("=", "+", "-", "@"):
            return f"'{value}"
    return value


def generate_xlsx_export(entries: Iterable[Entry]) -> io.BytesIO:
    """
    Generates a professionally styled XLSX file in memory:
    - Bold header with light fill
    - Frozen header pane
    - Real numeric cell values with 2-decimal formatting (#,##0.00)
    - Dates formatted as DD/MM/YYYY
    - Auto-fit column widths based on maximum content length
    - Formula injection protection on text fields
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Challan Entries"

    # Write header
    ws.append(EXPORT_COLUMNS)

    # Header styling
    header_font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    header_fill = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
    for col_idx in range(1, len(EXPORT_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(
            horizontal="center" if col_idx == 13 else "left",
            vertical="center"
        )

    # Freeze header row
    ws.freeze_panes = "A2"

    col_max_lengths = [len(h) for h in EXPORT_COLUMNS]

    for row_idx, entry in enumerate(entries, start=2):
        created_str = entry.created_at.strftime("%d/%m/%Y") if entry.created_at else ""

        qty_val = float(entry.quantity)
        unit_price_val = float(entry.unit_price)
        subtotal_val = float(entry.subtotal) if entry.subtotal is not None else float(entry.quantity * entry.unit_price)
        gst_type_val = (entry.gst_type or "none").upper()
        gst_rate_val = float(entry.gst_rate) if entry.gst_rate is not None else 0.0
        gst_amount_val = float(entry.gst_amount) if entry.gst_amount is not None else 0.0
        total_price_val = float(entry.total_price)

        row_values = [
            sanitize_formula_injection(entry.challan_no),
            sanitize_formula_injection(entry.vehicle_no),
            sanitize_formula_injection(entry.party_name or ""),
            sanitize_formula_injection(entry.product),
            sanitize_formula_injection(entry.destination),
            qty_val,
            unit_price_val,
            subtotal_val,
            gst_type_val,
            gst_rate_val if entry.gst_rate is not None else "",
            gst_amount_val,
            total_price_val,
            created_str,
        ]
        ws.append(row_values)

        # Apply specific numeric formatting
        # Quantity (col 6)
        qty_cell = ws.cell(row=row_idx, column=6)
        qty_cell.number_format = "#,##0.00" if (entry.quantity % 1 != 0) else "#,##0"
        qty_cell.alignment = Alignment(horizontal="right")

        # Unit Price (col 7)
        unit_price_cell = ws.cell(row=row_idx, column=7)
        unit_price_cell.number_format = "#,##0.00"
        unit_price_cell.alignment = Alignment(horizontal="right")

        # Subtotal (col 8)
        subtotal_cell = ws.cell(row=row_idx, column=8)
        subtotal_cell.number_format = "#,##0.00"
        subtotal_cell.alignment = Alignment(horizontal="right")

        # GST Type (col 9)
        gst_type_cell = ws.cell(row=row_idx, column=9)
        gst_type_cell.alignment = Alignment(horizontal="center")

        # GST Rate (col 10)
        gst_rate_cell = ws.cell(row=row_idx, column=10)
        if entry.gst_rate is not None:
            gst_rate_cell.number_format = "0.00"
            gst_rate_cell.alignment = Alignment(horizontal="right")

        # GST Amount (col 11)
        gst_amount_cell = ws.cell(row=row_idx, column=11)
        gst_amount_cell.number_format = "#,##0.00"
        gst_amount_cell.alignment = Alignment(horizontal="right")

        # Total Price (col 12)
        total_price_cell = ws.cell(row=row_idx, column=12)
        total_price_cell.number_format = "#,##0.00"
        total_price_cell.alignment = Alignment(horizontal="right")

        # Created At alignment (col 13)
        date_cell = ws.cell(row=row_idx, column=13)
        date_cell.alignment = Alignment(horizontal="center")

        # Track maximum length for auto column widths
        for col_idx, val in enumerate(row_values):
            val_str = f"{val:.2f}" if isinstance(val, float) else str(val)
            if len(val_str) > col_max_lengths[col_idx]:
                col_max_lengths[col_idx] = len(val_str)

    # Set calculated column widths with padding
    for col_idx, max_len in enumerate(col_max_lengths, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_csv_export(entries: Iterable[Entry]) -> io.BytesIO:
    """
    Generates a UTF-8 with BOM encoded CSV file in memory:
    - Standard comma delimiter with proper quoting
    - Prices formatted with 2 decimal places
    - Dates formatted as DD/MM/YYYY
    - Formula injection protection on text fields
    """
    output = io.BytesIO()
    # Write UTF-8 BOM so Excel opens it with proper encoding on Windows
    output.write(codecs.BOM_UTF8)

    text_stream = io.StringIO()
    writer = csv.writer(text_stream, quoting=csv.QUOTE_MINIMAL, lineterminator="\r\n")

    writer.writerow(EXPORT_COLUMNS)

    for entry in entries:
        created_str = entry.created_at.strftime("%d/%m/%Y") if entry.created_at else ""
        qty_str = f"{entry.quantity:.2f}" if (entry.quantity % 1 != 0) else f"{entry.quantity:.0f}"
        subtotal_val = entry.subtotal if entry.subtotal is not None else (entry.quantity * entry.unit_price)
        gst_rate_str = f"{entry.gst_rate:.2f}" if entry.gst_rate is not None else ""
        gst_amount_val = entry.gst_amount if entry.gst_amount is not None else 0.0

        row = [
            sanitize_formula_injection(entry.challan_no),
            sanitize_formula_injection(entry.vehicle_no),
            sanitize_formula_injection(entry.party_name or ""),
            sanitize_formula_injection(entry.product),
            sanitize_formula_injection(entry.destination),
            qty_str,
            f"{entry.unit_price:.2f}",
            f"{subtotal_val:.2f}",
            (entry.gst_type or "none").upper(),
            gst_rate_str,
            f"{gst_amount_val:.2f}",
            f"{entry.total_price:.2f}",
            created_str,
        ]
        writer.writerow(row)

    output.write(text_stream.getvalue().encode("utf-8"))
    output.seek(0)
    return output
