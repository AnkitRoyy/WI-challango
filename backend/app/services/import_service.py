import io
import csv
import math
import uuid
from datetime import datetime, timedelta, timezone
from decimal import Decimal, InvalidOperation
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import pandas as pd
from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from app.core.vehicle import normalize_vehicle_no, validate_indian_vehicle_format
from app.models.entry import Entry
from app.models.audit_log import AuditAction
from app.models.import_preview import ImportPreview
from app.schemas.import_schema import (
    ImportRowPreview,
    ImportPreviewResponse,
    ImportCommitResponse,
)
from app.services.entry_service import log_entry_audit, compute_entry_pricing
from app.services.export_service import sanitize_formula_injection

TEMPLATE_COLUMNS = [
    "Challan No",
    "Vehicle No",
    "Party",
    "Product",
    "Destination",
    "Quantity",
    "Unit Price",
    "Subtotal",
    "GST Type",
    "GST Rate (%)",
    "GST Amount",
    "Total Price",
]

REQUIRED_HEADERS_LOWER = {
    "challan no",
    "vehicle no",
    "product",
    "destination",
    "quantity",
    "unit price",
}


def generate_import_template() -> io.BytesIO:
    """
    Generates a blank Excel template (.xlsx) with exact required and optional headers,
    bold frozen header row, and one example row (italicized and grayed out)
    with a clear note indicating it should be deleted before uploading.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Challan Template"

    # Write header
    ws.append(TEMPLATE_COLUMNS)

    # Style header row
    header_font = Font(name="Calibri", size=11, bold=True, color="1F497D")
    header_fill = PatternFill(start_color="E9EEF4", end_color="E9EEF4", fill_type="solid")
    for col_idx in range(1, len(TEMPLATE_COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center" if col_idx in (2, 9) else "left", vertical="center")

    # Freeze header row
    ws.freeze_panes = "A2"

    # Example row (Row 2) - grayed out & italicized
    example_values = [
        "CH1001",
        "DL01AB1234",
        "Acme Enterprises",
        "Steel Rods 12mm",
        "Mumbai, Maharashtra",
        10,
        500.00,
        5000.00,
        "cgst_sgst",
        18.00,
        900.00,
        5900.00,
    ]
    ws.append(example_values)

    example_font = Font(name="Calibri", size=10, italic=True, color="7F7F7F")
    for col_idx in range(1, len(TEMPLATE_COLUMNS) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.font = example_font
        if col_idx in (6, 7, 8, 10, 11, 12):
            cell.number_format = "#,##0.00"

    # Add descriptive comment on cell A2
    comment = Comment(
        "NOTE: This row is just an example demonstrating the format.\nPlease delete or overwrite this row before uploading.",
        "ChallanGo System"
    )
    comment.width = 300
    comment.height = 70
    ws.cell(row=2, column=1).comment = comment

    # Set column widths
    column_widths = [18, 18, 22, 28, 24, 14, 14, 15, 14, 14, 15, 16]
    for col_idx, width in enumerate(column_widths, start=1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = width

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output


def generate_import_csv_template() -> io.BytesIO:
    """
    Generates a blank CSV template (.csv) with exact column headers including Party and GST fields,
    and one example row matching delivery challan requirements.
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(TEMPLATE_COLUMNS)
    writer.writerow([
        "CH1001",
        "DL01AB1234",
        "WEST INDUSTRIES",
        "Steel Rods 12mm",
        "Indore, Madhya Pradesh",
        10,
        500.00,
        5000.00,
        "cgst_sgst",
        18.00,
        900.00,
        5900.00,
    ])
    mem = io.BytesIO()
    mem.write(output.getvalue().encode("utf-8"))
    mem.seek(0)
    return mem



def _clean_str(val: Any) -> str:
    if val is None:
        return ""
    s = str(val).strip()
    if s.lower() in ("none", "nan", ""):
        return ""
    return s


def _clean_numeric_str(val: Any) -> Optional[str]:
    """Cleans numeric input strings (removes commas, spaces, currency symbols)."""
    cleaned = _clean_str(val)
    if not cleaned:
        return None
    s = cleaned.replace(",", "").replace("₹", "").replace("$", "").strip()
    return s if s else None


def parse_and_preview_import(
    file_bytes: bytes,
    filename: str,
    db: Session,
    user_id: int,
) -> ImportPreviewResponse:
    """
    Parses an uploaded .xlsx or .csv file and validates each row:
    - Validates headers
    - Skips fully blank rows
    - Detects in-file duplicates and database duplicates
    - Validates numeric quantity (> 0) and unit price (>= 0)
    - Recomputes total price and warns on discrepancies
    - Normalizes vehicle number and generates format warnings
    - Stores preview state in database with a 30-minute expiry
    Saves NOTHING to the entries table at this step.
    """
    is_csv = filename.lower().endswith(".csv")

    try:
        if is_csv:
            df = pd.read_csv(io.BytesIO(file_bytes), dtype=object, keep_default_na=False)
        else:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=False)
            ws = wb.active
            data = list(ws.iter_rows(values_only=True))
            if not data:
                raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Uploaded spreadsheet is empty")
            raw_headers = [str(c).strip() if c is not None else "" for c in data[0]]
            raw_rows = data[1:]
            df = pd.DataFrame(raw_rows, columns=raw_headers)
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Failed to parse spreadsheet file: {str(e)}",
        )

    # Check 10,000 row limit before heavy processing
    if len(df) > 10000:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File contains over 10,000 rows. Maximum allowed per import is 10,000 data rows",
        )

    # Normalize headers
    col_mapping: Dict[str, str] = {}
    for col in df.columns:
        cleaned_col = str(col).strip().lower()
        col_mapping[cleaned_col] = str(col)

    missing_headers = [h for h in REQUIRED_HEADERS_LOWER if h not in col_mapping]
    if missing_headers:
        display_missing = [h.title() for h in missing_headers]
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Missing required columns in file: {', '.join(display_missing)}",
        )

    # Map column names
    col_challan = col_mapping["challan no"]
    col_vehicle = col_mapping["vehicle no"]
    col_product = col_mapping["product"]
    col_destination = col_mapping["destination"]
    col_qty = col_mapping["quantity"]
    col_price = col_mapping["unit price"]
    col_party = col_mapping.get("party") or col_mapping.get("party name")
    col_subtotal = col_mapping.get("subtotal")
    col_gst_type = col_mapping.get("gst type")
    col_gst_rate = col_mapping.get("gst rate (%)") or col_mapping.get("gst rate")
    col_gst_amount = col_mapping.get("gst amount")
    col_total = col_mapping.get("total price")

    # Fetch active database entries for duplicate lookup (challan_no, product)
    active_db_entries = {
        (c.strip().lower(), p.strip().lower()): c
        for c, p in db.query(Entry.challan_no, Entry.product).filter(Entry.is_deleted.is_(False)).all()
    }

    seen_in_file: Dict[Tuple[str, str], int] = {}
    parsed_rows: List[Dict[str, Any]] = []

    ok_count = 0
    warning_count = 0
    duplicate_count = 0
    error_count = 0

    row_index = 0
    for idx, raw_row in df.iterrows():
        # Excel data row number (header is row 1)
        file_row_num = idx + 2

        raw_challan = _clean_str(raw_row.get(col_challan))
        raw_vehicle = _clean_str(raw_row.get(col_vehicle))
        raw_product = _clean_str(raw_row.get(col_product))
        raw_destination = _clean_str(raw_row.get(col_destination))
        raw_qty = _clean_str(raw_row.get(col_qty))
        raw_unit_price = _clean_str(raw_row.get(col_price))
        raw_party = _clean_str(raw_row.get(col_party)) if col_party else ""
        raw_gst_type = _clean_str(raw_row.get(col_gst_type)) if col_gst_type else ""
        raw_gst_rate = _clean_numeric_str(raw_row.get(col_gst_rate)) if col_gst_rate else None
        raw_subtotal = _clean_numeric_str(raw_row.get(col_subtotal)) if col_subtotal else None
        raw_gst_amount = _clean_numeric_str(raw_row.get(col_gst_amount)) if col_gst_amount else None
        raw_total_price = _clean_str(raw_row.get(col_total)) if col_total else None

        # Check if fully blank row
        if not any([raw_challan, raw_vehicle, raw_product, raw_destination, raw_qty, raw_unit_price, raw_party]):
            continue  # Skip fully blank rows silently

        row_index += 1
        row_errors: List[str] = []
        row_warnings: List[str] = []
        is_db_duplicate = False

        # Formula injection sanitization on text fields
        safe_challan = sanitize_formula_injection(raw_challan)
        safe_vehicle = sanitize_formula_injection(raw_vehicle)
        safe_product = sanitize_formula_injection(raw_product)
        safe_destination = sanitize_formula_injection(raw_destination)
        safe_party = sanitize_formula_injection(raw_party)

        # Presence checks
        if not safe_challan:
            row_errors.append("Challan No is required")
        if not safe_vehicle:
            row_errors.append("Vehicle No is required")
        if not safe_product:
            row_errors.append("Product is required")
        if not safe_destination:
            row_errors.append("Destination is required")

        # Parse Quantity
        qty_str = _clean_numeric_str(raw_qty)
        parsed_qty: Optional[Decimal] = None
        if qty_str is None or qty_str == "":
            row_errors.append("Quantity is required")
        else:
            try:
                parsed_qty = Decimal(qty_str)
                if parsed_qty <= Decimal("0"):
                    row_errors.append("Quantity must be greater than 0")
            except (InvalidOperation, ValueError):
                row_errors.append("Quantity is not a valid number")

        # Parse Unit Price
        unit_price_str = _clean_numeric_str(raw_unit_price)
        parsed_unit_price: Optional[Decimal] = None
        if unit_price_str is None or unit_price_str == "":
            row_errors.append("Unit Price is required")
        else:
            try:
                parsed_unit_price = Decimal(unit_price_str)
                if parsed_unit_price < Decimal("0"):
                    row_errors.append("Unit Price must be greater than or equal to 0")
            except (InvalidOperation, ValueError):
                row_errors.append("Unit Price is not a valid number")

        # Validate GST Type
        parsed_gst_type = "none"
        if raw_gst_type:
            gt_clean = raw_gst_type.lower()
            if gt_clean in ("none", "no gst", "0", ""):
                parsed_gst_type = "none"
            elif gt_clean in ("cgst_sgst", "cgst+sgst", "cgst and sgst", "cgst/sgst", "cgst_and_sgst", "cgst + sgst"):
                parsed_gst_type = "cgst_sgst"
            elif gt_clean in ("igst",):
                parsed_gst_type = "igst"
            else:
                row_errors.append(f"Invalid GST Type '{raw_gst_type}'. Must be 'none', 'cgst_sgst', or 'igst'")

        # Validate GST Rate
        parsed_gst_rate: Optional[Decimal] = None
        if raw_gst_rate is not None and raw_gst_rate != "":
            try:
                parsed_gst_rate = Decimal(raw_gst_rate)
                if parsed_gst_rate < Decimal("0") or parsed_gst_rate > Decimal("100"):
                    row_errors.append("GST Rate must be between 0 and 100")
                    parsed_gst_rate = None
            except (InvalidOperation, ValueError):
                row_errors.append("GST Rate is not a valid number")
                parsed_gst_rate = None

        # Calculate Pricing and Discrepancies
        computed_subtotal: Optional[Decimal] = None
        computed_gst_amount: Optional[Decimal] = None
        computed_total: Optional[Decimal] = None
        effective_gst_type = parsed_gst_type
        effective_gst_rate = parsed_gst_rate

        if parsed_qty is not None and parsed_qty > 0 and parsed_unit_price is not None and parsed_unit_price >= 0:
            computed_subtotal, effective_gst_type, effective_gst_rate, computed_gst_amount, computed_total = compute_entry_pricing(
                quantity=parsed_qty,
                unit_price=parsed_unit_price,
                gst_type=parsed_gst_type,
                gst_rate=parsed_gst_rate,
            )

            # Check provided Subtotal against computed Subtotal
            if raw_subtotal is not None and raw_subtotal != "":
                try:
                    provided_subtotal = Decimal(raw_subtotal)
                    if abs(provided_subtotal - computed_subtotal) > Decimal("0.05"):
                        row_warnings.append(
                            f"Subtotal does not match Quantity x Unit Price (expected {computed_subtotal:.2f}, got {provided_subtotal:.2f})"
                        )
                except (InvalidOperation, ValueError):
                    pass

            # Check provided Total Price against computed Total Price (Subtotal + GST Amount)
            total_price_str = _clean_numeric_str(raw_total_price)
            if total_price_str is not None and total_price_str != "":
                try:
                    provided_total = Decimal(total_price_str)
                    if abs(provided_total - computed_total) > Decimal("0.05"):
                        row_warnings.append(
                            f"Total Price does not match Subtotal + GST Amount (expected {computed_total:.2f}, got {provided_total:.2f})"
                        )
                except (InvalidOperation, ValueError):
                    row_warnings.append(f"Total Price is not a valid number; using computed value {computed_total:.2f}")

        # Vehicle normalization & format validation
        normalized_vehicle = normalize_vehicle_no(safe_vehicle)
        if normalized_vehicle:
            is_valid_v, v_warning = validate_indian_vehicle_format(normalized_vehicle)
            if not is_valid_v and v_warning:
                row_warnings.append(v_warning)

        # Duplicate checking on (Challan No, Product)
        if safe_challan and safe_product:
            pair_key = (safe_challan.lower(), safe_product.lower())
            if pair_key in seen_in_file:
                row_errors.append(
                    f"Duplicate (Challan No, Product) in file; previously seen on row {seen_in_file[pair_key]}"
                )
            else:
                seen_in_file[pair_key] = file_row_num
                if pair_key in active_db_entries:
                    is_db_duplicate = True
                    row_warnings.append(
                        f"Matches existing entry in database (Challan No '{safe_challan}', Product '{safe_product}')"
                    )

        # Determine row status
        if row_errors:
            row_status = "error"
            error_count += 1
            all_messages = row_errors + row_warnings
        elif is_db_duplicate:
            row_status = "duplicate"
            duplicate_count += 1
            all_messages = row_warnings
        elif row_warnings:
            row_status = "warning"
            warning_count += 1
            all_messages = row_warnings
        else:
            row_status = "ok"
            ok_count += 1
            all_messages = []

        row_data = {
            "challan_no": safe_challan,
            "vehicle_no": normalized_vehicle,
            "party_name": safe_party if safe_party else None,
            "product": safe_product,
            "destination": safe_destination,
            "quantity": str(parsed_qty) if parsed_qty is not None else raw_qty,
            "unit_price": str(parsed_unit_price) if parsed_unit_price is not None else raw_unit_price,
            "gst_type": effective_gst_type,
            "gst_rate": str(effective_gst_rate) if effective_gst_rate is not None else None,
            "subtotal": str(computed_subtotal) if computed_subtotal is not None else None,
            "gst_amount": str(computed_gst_amount) if computed_gst_amount is not None else None,
            "total_price": str(computed_total) if computed_total is not None else None,
            "provided_total_price": str(raw_total_price) if raw_total_price is not None else None,
        }

        parsed_rows.append({
            "row_number": file_row_num,
            "status": row_status,
            "messages": all_messages,
            "data": row_data,
        })

    preview_id = str(uuid.uuid4())
    now = datetime.now(timezone.utc)
    expires_at = now + timedelta(minutes=30)

    summary_dict = {
        "total_rows": len(parsed_rows),
        "ok_count": ok_count,
        "warning_count": warning_count,
        "duplicate_count": duplicate_count,
        "error_count": error_count,
    }

    # Store preview server-side
    preview_record = ImportPreview(
        id=preview_id,
        user_id=user_id,
        filename=filename,
        summary=summary_dict,
        rows=parsed_rows,
        created_at=now,
        expires_at=expires_at,
    )
    db.add(preview_record)
    db.commit()

    row_objects = [
        ImportRowPreview(
            row_number=r["row_number"],
            status=r["status"],
            messages=r["messages"],
            data=r["data"],
        )
        for r in parsed_rows
    ]

    return ImportPreviewResponse(
        preview_id=preview_id,
        filename=filename,
        total_rows=len(parsed_rows),
        ok_count=ok_count,
        warning_count=warning_count,
        duplicate_count=duplicate_count,
        error_count=error_count,
        rows=row_objects,
    )


def commit_import(
    preview_id: str,
    duplicate_strategy: str,
    db: Session,
    user_id: int,
) -> ImportCommitResponse:
    """
    Commits a validated preview into the database in a single atomic transaction:
    - Never inserts rows with status 'error'
    - Inserts rows with status 'ok' or 'warning'
    - For 'duplicate' rows:
      - 'skip': leaves existing DB row untouched, counts as skipped
      - 'update': updates existing DB row fields and recomputes total_price server-side
    - Writes an audit_log record with summary metrics
    - Invalidates the preview ID upon completion
    """
    now = datetime.now(timezone.utc)
    preview = db.query(ImportPreview).filter(ImportPreview.id == preview_id).first()
    if not preview:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Import preview not found. Please upload the file again",
        )

    if preview.expires_at < now:
        db.delete(preview)
        db.commit()
        raise HTTPException(
            status_code=status.HTTP_410_GONE,
            detail="Import preview has expired. Please upload the file again",
        )

    inserted = 0
    updated = 0
    skipped = 0
    failed = preview.summary.get("error_count", 0)

    try:
        for row in preview.rows:
            r_status = row.get("status")
            data = row.get("data", {})

            if r_status == "error":
                continue  # Error rows are never inserted

            qty = Decimal(str(data["quantity"]))
            unit_p = Decimal(str(data["unit_price"]))
            raw_gt = data.get("gst_type")
            raw_gr = Decimal(str(data["gst_rate"])) if data.get("gst_rate") is not None and str(data["gst_rate"]) != "" else None
            subtotal, eff_type, eff_rate, gst_amt, server_total = compute_entry_pricing(
                quantity=qty,
                unit_price=unit_p,
                gst_type=raw_gt,
                gst_rate=raw_gr,
            )

            dest = data.get("destination") or "Mumbai"
            party_name = data.get("party_name")

            if r_status in ("ok", "warning"):
                new_entry = Entry(
                    challan_no=data["challan_no"],
                    vehicle_no=data["vehicle_no"],
                    party_name=party_name,
                    product=data["product"],
                    destination=dest,
                    destination_lat=None,
                    destination_lng=None,
                    quantity=qty,
                    unit_price=unit_p,
                    gst_type=eff_type,
                    gst_rate=eff_rate,
                    subtotal=subtotal,
                    gst_amount=gst_amt,
                    total_price=server_total,
                    created_at=now,
                    updated_at=now,
                    created_by=user_id,
                    updated_by=user_id,
                    is_deleted=False,
                )
                db.add(new_entry)
                inserted += 1

            elif r_status == "duplicate":
                if duplicate_strategy == "skip":
                    skipped += 1
                elif duplicate_strategy == "update":
                    existing = (
                        db.query(Entry)
                        .filter(
                            Entry.challan_no == data["challan_no"],
                            Entry.product == data["product"],
                            Entry.is_deleted.is_(False),
                        )
                        .first()
                    )
                    if existing:
                        existing.vehicle_no = data["vehicle_no"]
                        existing.party_name = party_name
                        existing.product = data["product"]
                        if "destination" in data and data["destination"]:
                            existing.destination = data["destination"]
                        existing.quantity = qty
                        existing.unit_price = unit_p
                        existing.gst_type = eff_type
                        existing.gst_rate = eff_rate
                        existing.subtotal = subtotal
                        existing.gst_amount = gst_amt
                        existing.total_price = server_total
                        existing.updated_by = user_id
                        existing.updated_at = now
                        updated += 1
                    else:
                        # Fallback: if not found, insert
                        new_entry = Entry(
                            challan_no=data["challan_no"],
                            vehicle_no=data["vehicle_no"],
                            party_name=party_name,
                            product=data["product"],
                            destination=dest,
                            destination_lat=None,
                            destination_lng=None,
                            quantity=qty,
                            unit_price=unit_p,
                            gst_type=eff_type,
                            gst_rate=eff_rate,
                            subtotal=subtotal,
                            gst_amount=gst_amt,
                            total_price=server_total,
                            created_at=now,
                            updated_at=now,
                            created_by=user_id,
                            updated_by=user_id,
                            is_deleted=False,
                        )
                        db.add(new_entry)
                        inserted += 1

        # Write audit log for the whole import
        log_entry_audit(
            db=db,
            user_id=user_id,
            action=AuditAction.IMPORT,
            entry_id=None,
            details={
                "filename": preview.filename,
                "duplicate_strategy": duplicate_strategy,
                "inserted": inserted,
                "updated": updated,
                "skipped": skipped,
                "failed": failed,
            },
        )

        # Delete preview record so it cannot be committed twice
        db.delete(preview)
        db.commit()

    except Exception as e:
        db.rollback()
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f"Import transaction failed: {str(e)}",
        )

    msg = f"Import completed: {inserted} inserted, {updated} updated, {skipped} skipped, {failed} failed."
    return ImportCommitResponse(
        inserted=inserted,
        updated=updated,
        skipped=skipped,
        failed=failed,
        message=msg,
    )
